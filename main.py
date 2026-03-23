import asyncio
import io
import os
import re
import sqlite3
import ipaddress
import threading
import warnings
import secrets
from datetime import datetime, timezone
from typing import Any
from urllib.parse import quote_plus, urlparse

from flask import Flask, abort, g, make_response, redirect, render_template, request, url_for
from telegram import (
    Bot,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    InputFile,
    ReplyKeyboardMarkup,
    ReplyKeyboardRemove,
    Update,
)
from telegram.error import BadRequest
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    ConversationHandler,
    MessageHandler,
    filters,
)

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:
    Image = None
    ImageDraw = None
    ImageFont = None

warnings.filterwarnings(
    "ignore",
    message="If 'per_message=False', 'CallbackQueryHandler' will not be tracked*",
)

def load_local_env(path: str = ".env") -> None:
    if not os.path.exists(path):
        return
    with open(path, "r", encoding="utf-8-sig") as f:
        for raw in f:
            line = raw.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            key = key.strip()
            value = value.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = value


load_local_env()

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = (os.getenv("DB_PATH", "").strip() or os.path.join(BASE_DIR, "app.db"))
BOT_TOKEN = os.getenv("BOT_TOKEN", "TOKENNI_BU_YERGA_QOYING")
BASE_SITE_URL = os.getenv("BASE_SITE_URL", "http://localhost:5000")
TG_CHANNEL_URL = os.getenv("TG_CHANNEL_URL", "https://t.me/your_channel")
YOUTUBE_URL = os.getenv("YOUTUBE_URL", "https://youtube.com/@your_channel")
BOT_LINK = os.getenv("BOT_LINK", "https://t.me/your_bot")

ADMIN_IDS = {
    int(x.strip())
    for x in os.getenv("ADMIN_IDS", "").split(",")
    if x.strip().isdigit()
}
SUPER_ADMIN_ID = int(os.getenv("SUPER_ADMIN_ID", "0") or 0)
if SUPER_ADMIN_ID == 0 and ADMIN_IDS:
    SUPER_ADMIN_ID = sorted(ADMIN_IDS)[0]

REGISTER_FIRST, REGISTER_LAST, REGISTER_REGION = range(3)
ADMIN_ACTION, ADMIN_TEST_NUMBER, ADMIN_KEYS = range(3, 6)
IMG_TEST_TITLE, IMG_QUESTION_IMAGE, IMG_QUESTION_ANSWER, IMG_QUESTION_NEXT = range(6, 10)
EDIT_FIRST, EDIT_LAST, EDIT_REGION = range(10, 13)
START_CHOICE = 13
TEST_CODE = 14

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET", "change-me-please")


@app.context_processor
def inject_template_globals() -> dict[str, str]:
    return {"bot_link": BOT_LINK}


def is_valid_url(url: str) -> bool:
    return isinstance(url, str) and (url.startswith("http://") or url.startswith("https://"))


def is_button_safe_url(url: str) -> bool:
    if not is_valid_url(url):
        return False
    try:
        parsed = urlparse(url)
        host = (parsed.hostname or "").lower()
        if not host:
            return False
        if host in {"localhost", "127.0.0.1", "0.0.0.0"}:
            return False
        if parsed.scheme not in {"http", "https"}:
            return False
        try:
            ip = ipaddress.ip_address(host)
            if ip.is_loopback:
                return False
        except ValueError:
            pass
        return True
    except Exception:
        return False


def region_keyboard() -> ReplyKeyboardMarkup:
    keyboard = [
        ["Toshkent", "Samarqand", "Buxoro"],
        ["Andijon", "Namangan", "Farg'ona"],
        ["Navoiy", "Qashqadaryo", "Surxondaryo"],
        ["Xorazm", "Jizzax", "Sirdaryo"],
        ["Qoraqalpog'iston"],
    ]
    return ReplyKeyboardMarkup(keyboard, resize_keyboard=True, one_time_keyboard=True)


def start_menu_keyboard() -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(
        [["Test yaratish", "Testda qatnashish"]],
        resize_keyboard=True,
        one_time_keyboard=True,
    )


def image_test_next_keyboard() -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(
        [["Yana savol", "Bo'ldi"]],
        resize_keyboard=True,
        one_time_keyboard=True,
    )


def open_db_connection() -> sqlite3.Connection:
    db_path = (DB_PATH or "").strip() or os.path.join(BASE_DIR, "app.db")
    if not os.path.isabs(db_path):
        db_path = os.path.join(BASE_DIR, db_path)
    db_dir = os.path.dirname(db_path)
    if db_dir:
        os.makedirs(db_dir, exist_ok=True)

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    return conn


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        g.db = open_db_connection()
    return g.db
@app.teardown_appcontext
def close_db(_: Any) -> None:
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db() -> None:
    conn = open_db_connection()
    cur = conn.cursor()
    cur.executescript(
        """
        CREATE TABLE IF NOT EXISTS users (
            telegram_id INTEGER PRIMARY KEY,
            first_name TEXT NOT NULL,
            last_name TEXT NOT NULL,
            region TEXT NOT NULL,
            access_key TEXT NOT NULL,
            registered_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS tests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL UNIQUE,
            description TEXT,
            access_code TEXT
        );

        CREATE TABLE IF NOT EXISTS questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            test_id INTEGER NOT NULL,
            text TEXT NOT NULL,
            option_a TEXT NOT NULL,
            option_b TEXT NOT NULL,
            option_c TEXT NOT NULL,
            option_d TEXT NOT NULL,
            image_path TEXT,
            correct_option TEXT NOT NULL,
            FOREIGN KEY(test_id) REFERENCES tests(id)
        );

        CREATE TABLE IF NOT EXISTS submissions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            telegram_id INTEGER NOT NULL,
            test_id INTEGER NOT NULL,
            score INTEGER NOT NULL,
            total INTEGER NOT NULL,
            submitted_at TEXT NOT NULL,
            UNIQUE(telegram_id, test_id)
        );

        CREATE TABLE IF NOT EXISTS submission_answers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            submission_id INTEGER NOT NULL,
            question_id INTEGER NOT NULL,
            selected_option TEXT NOT NULL,
            is_correct INTEGER NOT NULL,
            FOREIGN KEY(submission_id) REFERENCES submissions(id)
        );

        CREATE TABLE IF NOT EXISTS admins (
            telegram_id INTEGER PRIMARY KEY,
            is_super INTEGER NOT NULL DEFAULT 0,
            added_at TEXT NOT NULL
        );
        """
    )

    # Ensure access_code column exists on existing databases.
    cols = {row[1] for row in cur.execute("PRAGMA table_info(tests)").fetchall()}
    if "access_code" not in cols:
        cur.execute("ALTER TABLE tests ADD COLUMN access_code TEXT")

    # Unique index for access codes.
    cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_tests_access_code ON tests(access_code)")

    # Backfill missing access codes on existing rows.
    ensure_test_codes(conn)

    # Ensure image_path column exists on questions table.
    qcols = {row[1] for row in cur.execute("PRAGMA table_info(questions)").fetchall()}
    if "image_path" not in qcols:
        cur.execute("ALTER TABLE questions ADD COLUMN image_path TEXT")

    # Seed admins from environment.
    now = datetime.now(timezone.utc).isoformat()
    if SUPER_ADMIN_ID:
        cur.execute(
            "INSERT OR IGNORE INTO admins (telegram_id, is_super, added_at) VALUES (?, ?, ?)",
            (SUPER_ADMIN_ID, 1, now),
        )
    for admin_id in ADMIN_IDS:
        is_super = 1 if admin_id == SUPER_ADMIN_ID else 0
        cur.execute(
            "INSERT OR IGNORE INTO admins (telegram_id, is_super, added_at) VALUES (?, ?, ?)",
            (admin_id, is_super, now),
        )

    cur.execute("SELECT COUNT(*) FROM tests")
    if cur.fetchone()[0] == 0:
        cur.execute(
            "INSERT INTO tests (title, description, access_code) VALUES (?, ?, ?)",
            ("41-test", "Matematika bo'yicha qisqa test", "DEMO41"),
        )
        test_id = cur.lastrowid
        cur.executemany(
            """
            INSERT INTO questions (test_id, text, option_a, option_b, option_c, option_d, correct_option)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (test_id, "2 + 4 = ?", "4", "5", "6", "7", "c"),
                (test_id, "3 + 4 = ?", "5", "6", "7", "8", "c"),
                (test_id, "12 - 5 = ?", "6", "7", "8", "9", "b"),
            ],
        )

    conn.commit()
    conn.close()


def db_conn() -> sqlite3.Connection:
    return open_db_connection()
def get_user_by_telegram_id(telegram_id: int) -> sqlite3.Row | None:
    with db_conn() as conn:
        return conn.execute("SELECT * FROM users WHERE telegram_id = ?", (telegram_id,)).fetchone()


def save_user(telegram_id: int, first_name: str, last_name: str, region: str) -> str:
    existing = get_user_by_telegram_id(telegram_id)
    if existing:
        return existing["access_key"]

    access_key = secrets.token_urlsafe(18)
    now = datetime.now(timezone.utc).isoformat()

    with db_conn() as conn:
        conn.execute(
            "INSERT INTO users (telegram_id, first_name, last_name, region, access_key, registered_at) VALUES (?, ?, ?, ?, ?, ?)",
            (telegram_id, first_name, last_name, region, access_key, now),
        )
        conn.commit()

    return access_key


def update_user_profile(telegram_id: int, first_name: str, last_name: str, region: str) -> bool:
    with db_conn() as conn:
        cur = conn.execute(
            "UPDATE users SET first_name = ?, last_name = ?, region = ? WHERE telegram_id = ?",
            (first_name, last_name, region, telegram_id),
        )
        conn.commit()
        return cur.rowcount > 0


def upsert_user_with_access_key(
    telegram_id: int,
    first_name: str,
    last_name: str,
    region: str,
    access_key: str,
) -> None:
    now = datetime.now(timezone.utc).isoformat()
    with db_conn() as conn:
        existing = conn.execute("SELECT telegram_id FROM users WHERE telegram_id = ?", (telegram_id,)).fetchone()
        if existing:
            conn.execute(
                "UPDATE users SET first_name = ?, last_name = ?, region = ?, access_key = ? WHERE telegram_id = ?",
                (first_name, last_name, region, access_key, telegram_id),
            )
        else:
            conn.execute(
                "INSERT INTO users (telegram_id, first_name, last_name, region, access_key, registered_at) VALUES (?, ?, ?, ?, ?, ?)",
                (telegram_id, first_name, last_name, region, access_key, now),
            )
        conn.commit()


def build_site_link(
    telegram_id: int,
    access_key: str,
    first_name: str | None = None,
    last_name: str | None = None,
    region: str | None = None,
) -> str:
    params = [f"tg_id={telegram_id}", f"key={quote_plus(access_key)}"]
    if first_name:
        params.append(f"fn={quote_plus(first_name)}")
    if last_name:
        params.append(f"ln={quote_plus(last_name)}")
    if region:
        params.append(f"rg={quote_plus(region)}")
    return f"{BASE_SITE_URL}/?{'&'.join(params)}"


def is_valid_user(telegram_id: int, key: str) -> sqlite3.Row | None:
    try:
        with db_conn() as conn:
            return conn.execute(
                "SELECT * FROM users WHERE telegram_id = ? AND access_key = ?",
                (telegram_id, key),
            ).fetchone()
    except sqlite3.OperationalError as e:
        if "no such table" not in str(e).lower():
            raise
        init_db()
        with db_conn() as conn:
            return conn.execute(
                "SELECT * FROM users WHERE telegram_id = ? AND access_key = ?",
                (telegram_id, key),
            ).fetchone()


def is_admin(telegram_id: int) -> bool:
    if telegram_id in ADMIN_IDS or telegram_id == SUPER_ADMIN_ID:
        return True
    try:
        with db_conn() as conn:
            row = conn.execute(
                "SELECT 1 FROM admins WHERE telegram_id = ?",
                (telegram_id,),
            ).fetchone()
        return row is not None
    except sqlite3.OperationalError:
        return False


def is_super_admin(telegram_id: int) -> bool:
    if telegram_id == SUPER_ADMIN_ID and SUPER_ADMIN_ID != 0:
        return True
    try:
        with db_conn() as conn:
            row = conn.execute(
                "SELECT is_super FROM admins WHERE telegram_id = ?",
                (telegram_id,),
            ).fetchone()
        return bool(row and row["is_super"])
    except sqlite3.OperationalError:
        return False


def add_admin(telegram_id: int) -> bool:
    now = datetime.now(timezone.utc).isoformat()
    with db_conn() as conn:
        conn.execute(
            "INSERT OR IGNORE INTO admins (telegram_id, is_super, added_at) VALUES (?, ?, ?)",
            (telegram_id, 0, now),
        )
        conn.commit()
        row = conn.execute(
            "SELECT 1 FROM admins WHERE telegram_id = ?",
            (telegram_id,),
        ).fetchone()
        return row is not None


def remove_admin(telegram_id: int) -> bool:
    if telegram_id == SUPER_ADMIN_ID:
        return False
    with db_conn() as conn:
        cur = conn.execute(
            "DELETE FROM admins WHERE telegram_id = ? AND is_super = 0",
            (telegram_id,),
        )
        conn.commit()
        return cur.rowcount > 0


def list_admins() -> list[sqlite3.Row]:
    with db_conn() as conn:
        return conn.execute(
            "SELECT telegram_id, is_super, added_at FROM admins ORDER BY is_super DESC, telegram_id ASC"
        ).fetchall()


def generate_test_code(conn: sqlite3.Connection) -> str:
    # Short human-friendly code, uppercase letters + digits.
    alphabet = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
    while True:
        code = "".join(secrets.choice(alphabet) for _ in range(6))
        exists = conn.execute("SELECT 1 FROM tests WHERE access_code = ?", (code,)).fetchone()
        if not exists:
            return code


def ensure_test_codes(conn: sqlite3.Connection) -> None:
    rows = conn.execute("SELECT id FROM tests WHERE access_code IS NULL OR access_code = ''").fetchall()
    for r in rows:
        code = generate_test_code(conn)
        conn.execute("UPDATE tests SET access_code = ? WHERE id = ?", (code, r["id"]))


async def save_telegram_photo(update: Update, test_id: int, index: int) -> str | None:
    if not update.message or not update.message.photo:
        return None
    photo = update.message.photo[-1]
    file = await photo.get_file()
    uploads_dir = os.path.join(BASE_DIR, "static", "uploads")
    os.makedirs(uploads_dir, exist_ok=True)
    ts = int(datetime.now(timezone.utc).timestamp())
    filename = f"test_{test_id}_{index}_{ts}.jpg"
    full_path = os.path.join(uploads_dir, filename)
    await file.download_to_drive(full_path)
    return f"/static/uploads/{filename}"


def create_test_from_keys(test_number: str, keys: str) -> tuple[bool, str]:
    title = f"{test_number}-test"
    cleaned = keys.strip().lower()

    if not re.fullmatch(r"[abcd]+", cleaned):
        return False, "Kalitlar faqat a, b, c, d harflaridan iborat bo'lishi kerak."

    with db_conn() as conn:
        exists = conn.execute("SELECT id FROM tests WHERE title = ?", (title,)).fetchone()
        if exists:
            return False, f"{title} allaqachon mavjud."

        access_code = generate_test_code(conn)
        cur = conn.execute(
            "INSERT INTO tests (title, description, access_code) VALUES (?, ?, ?)",
            (title, f"Admin yaratgan kalitli test. Savollar soni: {len(cleaned)}", access_code),
        )
        test_id = cur.lastrowid
        rows = []
        for i, letter in enumerate(cleaned, start=1):
            rows.append(
                (
                    test_id,
                    f"{i}-savol: To'g'ri javobni belgilang",
                    "A variant",
                    "B variant",
                    "C variant",
                    "D variant",
                    None,
                    letter,
                )
            )
        conn.executemany(
            "INSERT INTO questions (test_id, text, option_a, option_b, option_c, option_d, image_path, correct_option) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            rows,
        )
        conn.commit()

        # Quick trace for production debugging of DB consistency issues.
        print(f"[ADMIN CREATE] db={DB_PATH} title={title} questions={len(cleaned)} code={access_code}")

    return True, f"✅ {title} yaratildi. Savollar soni: {len(cleaned)} ta. Test kodi: {access_code}"


def delete_test_by_number(test_number: str) -> tuple[bool, str]:
    title = f"{test_number}-test"
    with db_conn() as conn:
        test = conn.execute("SELECT id FROM tests WHERE title = ?", (title,)).fetchone()
        if not test:
            return False, f"{title} topilmadi."

        test_id = test["id"]
        submission_ids = conn.execute("SELECT id FROM submissions WHERE test_id = ?", (test_id,)).fetchall()
        sub_ids = [row["id"] for row in submission_ids]

        if sub_ids:
            placeholders = ",".join(["?"] * len(sub_ids))
            conn.execute(
                f"DELETE FROM submission_answers WHERE submission_id IN ({placeholders})",
                sub_ids,
            )

        conn.execute("DELETE FROM submissions WHERE test_id = ?", (test_id,))
        conn.execute("DELETE FROM questions WHERE test_id = ?", (test_id,))
        conn.execute("DELETE FROM tests WHERE id = ?", (test_id,))
        conn.commit()

        print(f"[ADMIN DELETE] db={DB_PATH} title={title} test_id={test_id}")

    return True, f"{title} o'chirildi."


def get_test_by_access_code(code: str) -> sqlite3.Row | None:
    if not code:
        return None
    with db_conn() as conn:
        return conn.execute("SELECT * FROM tests WHERE access_code = ?", (code,)).fetchone()

def get_test_results_rows(test_number: str) -> tuple[str | None, list[sqlite3.Row]]:
    title = f"{test_number}-test"
    with db_conn() as conn:
        test = conn.execute("SELECT id FROM tests WHERE title = ?", (title,)).fetchone()
        if not test:
            return None, []

        rows = conn.execute(
            """
            SELECT
                u.telegram_id,
                u.first_name,
                u.last_name,
                u.region,
                s.score,
                s.total,
                s.submitted_at
            FROM submissions s
            JOIN users u ON u.telegram_id = s.telegram_id
            WHERE s.test_id = ?
            ORDER BY s.submitted_at DESC
            """,
            (test["id"],),
        ).fetchall()

    return title, rows


def get_test_results_text(test_number: str) -> str:
    title, rows = get_test_results_rows(test_number)
    if not title:
        return f"❗ {test_number}-test topilmadi."

    if not rows:
        return f"ℹ️ {title} uchun hali natija yo'q."

    lines = [f"📊 {title} natijalari:"]
    for i, r in enumerate(rows, start=1):
        lines.append(f"{i}. {r['last_name']} {r['first_name']} - {r['score']}/{r['total']}")
    return "\n".join(lines)


def build_results_excel(test_number: str) -> tuple[io.BytesIO | None, str]:
    if Workbook is None:
        return None, "❗ Excel yuborish uchun openpyxl o'rnatilmagan. 'pip install openpyxl' qiling."

    title, rows = get_test_results_rows(test_number)
    if not title:
        return None, f"❗ {test_number}-test topilmadi."

    if not rows:
        return None, f"ℹ️ {title} uchun hali natija yo'q."

    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"

    ws.append(["#", "Ism", "Familiya", "Viloyat", "Telegram ID", "Natija", "Sana"])

    for i, r in enumerate(rows, start=1):
        ws.append(
            [
                i,
                r["first_name"],
                r["last_name"],
                r["region"],
                r["telegram_id"],
                f"{r['score']}/{r['total']}",
                r["submitted_at"],
            ]
        )

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream, f"{title}-natijalar.xlsx"

def certificate_text(telegram_id: int) -> str:
    with db_conn() as conn:
        user = conn.execute("SELECT first_name, last_name FROM users WHERE telegram_id = ?", (telegram_id,)).fetchone()
        stats = conn.execute(
            "SELECT COUNT(*) as cnt, COALESCE(SUM(score), 0) as total_score, COALESCE(SUM(total), 0) as total_q FROM submissions WHERE telegram_id = ?",
            (telegram_id,),
        ).fetchone()

    if not user or stats["cnt"] == 0:
        return ""

    percent = round((stats["total_score"] / stats["total_q"]) * 100, 1) if stats["total_q"] else 0
    code = f"CERT-{telegram_id}-{stats['cnt']}"
    return (
        f"SERTIFIKAT\n"
        f"Ism: {user['first_name']} {user['last_name']}\n"
        f"Telegram ID: {telegram_id}\n"
        f"Topshirilgan testlar: {stats['cnt']} ta\n"
        f"Umumiy natija: {percent}%\n"
        f"Sertifikat kodi: {code}"
    )




def generate_certificate(user, percent, tg_id, test_code, result_lines=None):
    if Image is None or ImageDraw is None or ImageFont is None:
        return None

    width = 1600
    height = 1000

    img = Image.new("RGB", (width, height), "#eeeeee")
    draw = ImageDraw.Draw(img)

    draw.rectangle([10, 10, width - 10, height - 10], outline="#2c3e50", width=40)
    draw.rectangle([80, 80, width - 80, height - 80], outline="#f1c40f", width=8)


    def load_font(size: int, bold: bool = False):
        font_candidates = [
            os.path.join("fonts", "DejaVuSans-Bold.ttf") if bold else os.path.join("fonts", "DejaVuSans.ttf"),
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        ]

        for fp in font_candidates:
            if fp and os.path.exists(fp):
                try:
                    return ImageFont.truetype(fp, size)
                except Exception:
                    pass

        return ImageFont.load_default()

    title_font = load_font(95, bold=True)
    subtitle_font = load_font(40)
    name_font = load_font(85, bold=True)
    text_font = load_font(48)
    small_font = load_font(36)

    def center(text, y, font, color="#2c3e50"):
        w = draw.textlength(text, font=font)
        draw.text(((width - w) / 2, y), text, fill=color, font=font)

    center("SERTIFIKAT", 170, title_font)
    center("Ushbu sertifikat matematika fanidan olingan bilim darajasini tasdiqlaydi", 300, subtitle_font)

    fullname = f"{user['surname'].upper()} {user['name'].upper()}"
    center(fullname, 430, name_font, "#e74c3c")

    line1 = "Matematika fanidan o'tkazilgan test sinovida"
    percent_text = int(percent) if float(percent).is_integer() else percent
    line2 = f"ishtirok etib, {percent_text}% natija qayd etdi."
    center(line1, 610, text_font)
    center(line2, 680, text_font)

    today = datetime.now().strftime("%d.%m.%Y")
    bottom_y = height - 140
    draw.text((140, bottom_y), f"Sana: {today}", fill="#2c3e50", font=small_font)

    academy = "Matematika Prime Akademiyasi"
    w = draw.textlength(academy, font=small_font)
    draw.text((width - w - 140, bottom_y), academy, fill="#2c3e50", font=small_font)

    safe_code = re.sub(r"[^a-zA-Z0-9_-]", "_", str(test_code))
    filename = f"cert_{tg_id}_{safe_code}.jpg"
    img.save(filename, quality=95)
    return filename

async def send_test_result_certificate(
    telegram_id: int,
    user_first_name: str,
    user_last_name: str,
    test_title: str,
    score: int,
    total: int,
    result_lines: list[str],
) -> None:
    if not BOT_TOKEN or BOT_TOKEN == "TOKENNI_BU_YERGA_QOYING":
        return

    percent = round((score / total) * 100, 1) if total else 0
    cert_path = generate_certificate(
        user={"name": user_first_name, "surname": user_last_name},
        percent=percent,
        tg_id=telegram_id,
        test_code=test_title,
        result_lines=result_lines,
    )
    if not cert_path:
        return

    bot = Bot(token=BOT_TOKEN)

    result_text = "\n".join(result_lines)
    await bot.send_message(
        chat_id=telegram_id,
        text=(
            f"📊 {test_title} natijalari\n"
            f"Natija: {score}/{total} ({percent}%)\n\n"
            f"{result_text}"
        ),
    )

    try:
        with open(cert_path, "rb") as f:
            await bot.send_document(
                chat_id=telegram_id,
                document=InputFile(f, filename=os.path.basename(cert_path)),
            )
    finally:
        try:
            os.remove(cert_path)
        except OSError:
            pass

def push_test_result_certificate(
    telegram_id: int,
    user_first_name: str,
    user_last_name: str,
    test_title: str,
    score: int,
    total: int,
    result_lines: list[str],
) -> None:
    def runner() -> None:
        try:
            asyncio.run(
                send_test_result_certificate(
                    telegram_id=telegram_id,
                    user_first_name=user_first_name,
                    user_last_name=user_last_name,
                    test_title=test_title,
                    score=score,
                    total=total,
                    result_lines=result_lines,
                )
            )
        except Exception as e:
            print(f"[CERT ERROR] {e}")

    threading.Thread(target=runner, daemon=True).start()

@app.route("/")
def test_list() -> str:
    tg_id = request.args.get("tg_id", type=int)
    key = request.args.get("key", default="", type=str)
    test_code = request.args.get("tc", default="", type=str).strip().upper()
    first_name = request.args.get("fn", default="", type=str).strip()
    last_name = request.args.get("ln", default="", type=str).strip()
    region = request.args.get("rg", default="", type=str).strip()
    search = request.args.get("q", default="", type=str).strip()

    if not tg_id or not key:
        return render_template("public_home.html", user=None, tg_channel=TG_CHANNEL_URL, youtube=YOUTUBE_URL)

    user = is_valid_user(tg_id, key)
    if not user and first_name and last_name and region:
        upsert_user_with_access_key(
            telegram_id=tg_id,
            first_name=first_name,
            last_name=last_name,
            region=region,
            access_key=key,
        )
        user = is_valid_user(tg_id, key)
    if not user:
        return render_template("public_home.html", user=None, tg_channel=TG_CHANNEL_URL, youtube=YOUTUBE_URL)

    db = get_db()
    tests = []
    code_error = ""
    if not test_code:
        code_error = "Testga kirish uchun botdan olingan test kodini kiriting."
    else:
        if search:
            tests = db.execute(
                "SELECT id, title, description, access_code FROM tests WHERE access_code = ? AND title LIKE ? ORDER BY id DESC",
                (test_code, f"%{search}%"),
            ).fetchall()
        else:
            tests = db.execute(
                "SELECT id, title, description, access_code FROM tests WHERE access_code = ? ORDER BY id DESC",
                (test_code,),
            ).fetchall()
        if not tests:
            code_error = "Test kodi noto'g'ri yoki test topilmadi."

    response = make_response(
        render_template(
            "tests.html",
            tests=tests,
            tg_id=tg_id,
            key=key,
            tc=test_code,
            search=search,
            code_error=code_error,
            user=user,
            tg_channel=TG_CHANNEL_URL,
            youtube=YOUTUBE_URL,
        )
    )
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


@app.route("/profile", methods=["GET", "POST"])
def profile_page() -> str:
    tg_id = request.args.get("tg_id", type=int)
    key = request.args.get("key", default="", type=str)

    if not tg_id or not key:
        return redirect(url_for("test_list"))

    user = is_valid_user(tg_id, key)
    if not user:
        return redirect(url_for("test_list"))

    if request.method == "POST":
        first_name = request.form.get("first_name", "").strip()
        last_name = request.form.get("last_name", "").strip()
        region = request.form.get("region", "").strip()

        if len(first_name) < 2 or len(last_name) < 2 or len(region) < 2:
            return render_template(
                "profile.html",
                user=user,
                tg_id=tg_id,
                key=key,
                error="Ma'lumotlarni to'g'ri kiriting.",
                tg_channel=TG_CHANNEL_URL,
                youtube=YOUTUBE_URL,
            )

        update_user_profile(tg_id, first_name, last_name, region)
        user = is_valid_user(tg_id, key)
        return render_template(
            "profile.html",
            user=user,
            tg_id=tg_id,
            key=key,
            success="Profil muvaffaqiyatli yangilandi.",
            tg_channel=TG_CHANNEL_URL,
            youtube=YOUTUBE_URL,
        )

    return render_template(
        "profile.html",
        user=user,
        tg_id=tg_id,
        key=key,
        tg_channel=TG_CHANNEL_URL,
        youtube=YOUTUBE_URL,
    )


@app.route("/test/<int:test_id>", methods=["GET", "POST"])
def test_detail(test_id: int) -> str:
    tg_id = request.args.get("tg_id", type=int)
    key = request.args.get("key", default="", type=str)
    test_code = request.args.get("tc", default="", type=str).strip().upper()

    if not tg_id or not key:
        return redirect(url_for("test_list"))

    user = is_valid_user(tg_id, key)
    if not user:
        return redirect(url_for("test_list"))

    db = get_db()
    test = db.execute("SELECT * FROM tests WHERE id = ?", (test_id,)).fetchone()
    if not test:
        abort(404)
    if not test_code or test_code != (test["access_code"] or "").upper():
        return redirect(url_for("test_list", tg_id=tg_id, key=key, tc=test_code))

    questions = db.execute("SELECT * FROM questions WHERE test_id = ? ORDER BY id", (test_id,)).fetchall()

    already = db.execute(
        "SELECT * FROM submissions WHERE telegram_id = ? AND test_id = ?",
        (tg_id, test_id),
    ).fetchone()

    if request.method == "POST":
        if already:
            return render_template(
                "test_detail.html",
                test=test,
                questions=questions,
                tg_id=tg_id,
                key=key,
                tc=test_code,
                already=already,
                user=user,
                tg_channel=TG_CHANNEL_URL,
                youtube=YOUTUBE_URL,
            )

        selected: dict[int, str] = {}
        for q in questions:
            val = request.form.get(f"q_{q['id']}", "").strip().lower()
            if val not in {"a", "b", "c", "d"}:
                return render_template(
                    "test_detail.html",
                    test=test,
                    questions=questions,
                    tg_id=tg_id,
                    key=key,
                    tc=test_code,
                    user=user,
                    error="Barcha savollarga javob berish majburiy.",
                    tg_channel=TG_CHANNEL_URL,
                    youtube=YOUTUBE_URL,
                )
            selected[q["id"]] = val

        score = 0
        results = []
        for q in questions:
            choice = selected[q["id"]]
            ok = int(choice == q["correct_option"])
            score += ok
            results.append({"question": q, "selected": choice, "is_correct": ok})

        now = datetime.now(timezone.utc).isoformat()
        cur = db.execute(
            "INSERT INTO submissions (telegram_id, test_id, score, total, submitted_at) VALUES (?, ?, ?, ?, ?)",
            (tg_id, test_id, score, len(questions), now),
        )
        sub_id = cur.lastrowid

        db.executemany(
            "INSERT INTO submission_answers (submission_id, question_id, selected_option, is_correct) VALUES (?, ?, ?, ?)",
            [(sub_id, item["question"]["id"], item["selected"], item["is_correct"]) for item in results],
        )
        db.commit()

        result_lines = []
        for idx, item in enumerate(results, start=1):
            selected_letter = item["selected"].upper()
            mark = "✅" if item["is_correct"] else "❌"
            result_lines.append(f"{idx}. {selected_letter} {mark}")

        push_test_result_certificate(
            telegram_id=tg_id,
            user_first_name=user["first_name"],
            user_last_name=user["last_name"],
            test_title=test["title"],
            score=score,
            total=len(questions),
            result_lines=result_lines,
        )

        return render_template(
            "test_result.html",
            test=test,
            results=results,
            score=score,
            total=len(questions),
            tg_id=tg_id,
            key=key,
            user=user,
            tg_channel=TG_CHANNEL_URL,
            youtube=YOUTUBE_URL,
        )

    return render_template(
        "test_detail.html",
        test=test,
        questions=questions,
        tg_id=tg_id,
        key=key,
        tc=test_code,
        already=already,
        user=user,
        tg_channel=TG_CHANNEL_URL,
        youtube=YOUTUBE_URL,
    )


async def _send_access_button_or_text(message, text: str, site_link: str) -> None:
    if is_button_safe_url(site_link):
        try:
            await message.reply_text(
                text,
                reply_markup=InlineKeyboardMarkup(
                    [
                        [InlineKeyboardButton("🔗 Kirish", url=site_link)],
                        [InlineKeyboardButton("👤 Profilni o'zgartirish", callback_data="edit_profile")],
                    ]
                ),
            )
            return
        except BadRequest:
            pass
    await message.reply_text(f"{text}\nKirish havolasi: {site_link}")


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data.clear()

    tg_id = update.effective_user.id
    existing = get_user_by_telegram_id(tg_id)

    if existing:
        await update.message.reply_text(
            "Test yaratasizmi yoki qatnashasizmi?",
            reply_markup=start_menu_keyboard(),
        )
        return START_CHOICE

    await update.message.reply_text("Salom! 👋\nRo'yxatdan o'tishni boshlaymiz.\nIsmingizni kiriting:")
    return REGISTER_FIRST


async def first_name(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    text = (update.message.text or "").strip()
    if len(text) < 2:
        await update.message.reply_text("Ism kamida 2 ta harfdan iborat bo'lishi kerak. Qayta kiriting 👇")
        return REGISTER_FIRST
    context.user_data["first_name"] = text
    await update.message.reply_text("Familiyangizni kiriting 👇")
    return REGISTER_LAST


async def last_name(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    text = (update.message.text or "").strip()
    if len(text) < 2:
        await update.message.reply_text("Familiya kamida 2 ta harfdan iborat bo'lishi kerak. Qayta kiriting 👇")
        return REGISTER_LAST
    context.user_data["last_name"] = text
    await update.message.reply_text("Viloyatingizni kiriting yoki pastdagi tugmadan tanlang 🗺️", reply_markup=region_keyboard())
    return REGISTER_REGION


async def region(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    text = (update.message.text or "").strip()
    if len(text) < 2:
        await update.message.reply_text("Viloyat nomini to'g'ri kiriting 👇")
        return REGISTER_REGION

    context.user_data["region"] = text
    tg_id = update.effective_user.id
    key = save_user(tg_id, context.user_data["first_name"], context.user_data["last_name"], context.user_data["region"])
    await update.message.reply_text(
        "✅ Ro'yxatdan muvaffaqiyatli o'tdingiz.\n"
        "Test yaratasizmi yoki qatnashasizmi?",
        reply_markup=start_menu_keyboard(),
    )
    return START_CHOICE


async def test_code(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    code = (update.message.text or "").strip().upper()
    if len(code) < 4:
        await update.message.reply_text("Test kodini to'g'ri kiriting.")
        return TEST_CODE

    test = get_test_by_access_code(code)
    if not test:
        await update.message.reply_text("Test kodi topilmadi. Qayta kiriting.")
        return TEST_CODE

    tg_id = update.effective_user.id
    user = get_user_by_telegram_id(tg_id)
    if not user:
        await update.message.reply_text("Avval /start orqali ro'yxatdan o'ting.")
        return ConversationHandler.END

    site_link = (
        f"{build_site_link(tg_id, user['access_key'], user['first_name'], user['last_name'], user['region'])}"
        f"&q={quote_plus(test['title'])}&tc={quote_plus(code)}&v={int(datetime.now().timestamp())}"
    )
    await _send_access_button_or_text(update.message, "✅ Testga kirish uchun havola:", site_link)
    return ConversationHandler.END


async def start_choice(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    text = (update.message.text or "").strip().lower()
    if "yarat" in text:
        await update.message.reply_text(
            "Test yaratish uchun @Murojatuchunadminbek shu adminga bog'laning !!!",
            reply_markup=ReplyKeyboardRemove(),
        )
        return ConversationHandler.END
    if "qatnash" in text:
        await update.message.reply_text(
            "Tayyor testda qatnashish uchun test kodini kiriting:",
            reply_markup=ReplyKeyboardRemove(),
        )
        return TEST_CODE

    await update.message.reply_text(
        "Iltimos, tugmadan tanlang: Test yaratish yoki Testda qatnashish.",
        reply_markup=start_menu_keyboard(),
    )
    return START_CHOICE


async def editprofile(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    tg_id = update.effective_user.id
    user = get_user_by_telegram_id(tg_id)
    if not user:
        await update.message.reply_text("Siz hali ro'yxatdan o'tmagansiz. /start ni bosing.")
        return ConversationHandler.END

    await update.message.reply_text(
        f"Joriy ism: {user['first_name']}\nYangi ismingizni kiriting:"
    )
    return EDIT_FIRST


async def editprofile_callback(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    query = update.callback_query
    await query.answer()

    tg_id = update.effective_user.id
    user = get_user_by_telegram_id(tg_id)
    if not user:
        await query.message.reply_text("Siz hali ro'yxatdan o'tmagansiz. /start ni bosing.")
        return ConversationHandler.END

    await query.message.reply_text(
        f"Joriy ism: {user['first_name']}\nYangi ismingizni kiriting:"
    )
    return EDIT_FIRST

async def edit_first(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    text = (update.message.text or "").strip()
    if len(text) < 2:
        await update.message.reply_text("Ism kamida 2 ta harf bo'lsin. Qayta kiriting:")
        return EDIT_FIRST
    context.user_data["edit_first_name"] = text
    await update.message.reply_text("Yangi familiyangizni kiriting:")
    return EDIT_LAST


async def edit_last(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    text = (update.message.text or "").strip()
    if len(text) < 2:
        await update.message.reply_text("Familiya kamida 2 ta harf bo'lsin. Qayta kiriting:")
        return EDIT_LAST
    context.user_data["edit_last_name"] = text
    await update.message.reply_text("Yangi viloyatni kiriting:", reply_markup=region_keyboard())
    return EDIT_REGION


async def edit_region(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    text = (update.message.text or "").strip()
    if len(text) < 2:
        await update.message.reply_text("Viloyat nomini to'g'ri kiriting:")
        return EDIT_REGION

    tg_id = update.effective_user.id
    ok = update_user_profile(
        tg_id,
        context.user_data.get("edit_first_name", ""),
        context.user_data.get("edit_last_name", ""),
        text,
    )

    if not ok:
        await update.message.reply_text("Profil yangilanmadi. /start ni bosing.", reply_markup=ReplyKeyboardRemove())
        return ConversationHandler.END

    user = get_user_by_telegram_id(tg_id)
    site_link = build_site_link(tg_id, user["access_key"], user["first_name"], user["last_name"], user["region"])
    await _send_access_button_or_text(update.message, "✅ Profil muvaffaqiyatli yangilandi.", site_link)
    return ConversationHandler.END


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data.clear()
    await update.message.reply_text(
        "Ro'yxatdan o'tish bekor qilindi. Qayta boshlash uchun /start ni bosing.",
        reply_markup=ReplyKeyboardRemove(),
    )
    return ConversationHandler.END


def admin_keyboard() -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(
        [["🧪 Test yaratish", "🖼 Rasmli test yaratish"], ["📊 Natijalarni olish", "🗑 Test o'chirish"], ["❌ Admin paneldan chiqish"]],
        resize_keyboard=True,
    )


async def admin_panel(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    tg_id = update.effective_user.id
    if not is_admin(tg_id):
        await update.message.reply_text("Kechirasiz, sizda admin huquqi yo'q ⛔")
        return ConversationHandler.END

    context.user_data["admin_action"] = None
    context.user_data["admin_state"] = "action"
    await update.message.reply_text("Admin panelga xush kelibsiz 👨‍💼", reply_markup=admin_keyboard())
    return ADMIN_ACTION


async def admin_action(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    text = (update.message.text or "").strip().lower()

    if text in {"🧪 test yaratish", "test yaratish"}:
        context.user_data["admin_action"] = "create"
        context.user_data["admin_state"] = "create_number"
        await update.message.reply_text("Test raqamini kiriting (masalan: 55) 🔢")
        return ADMIN_TEST_NUMBER

    if text in {"🖼 rasmli test yaratish", "rasmli test yaratish"}:
        context.user_data["admin_action"] = "image_create"
        context.user_data["admin_state"] = "image_title"
        await update.message.reply_text("Rasmli testga nom bering:")
        return IMG_TEST_TITLE

    if text in {"📊 natijalarni olish", "natijalarni olish"}:
        context.user_data["admin_action"] = "results"
        context.user_data["admin_state"] = "results_number"
        await update.message.reply_text("Qaysi test raqami natijasi kerak? (masalan: 55) 📈")
        return ADMIN_TEST_NUMBER

    if text in {"🗑 test o'chirish", "test o'chirish"}:
        context.user_data["admin_action"] = "delete"
        context.user_data["admin_state"] = "delete_number"
        await update.message.reply_text("Qaysi testni o'chiramiz? Raqamini kiriting (masalan: 12) 🗑")
        return ADMIN_TEST_NUMBER

    if text in {"❌ admin paneldan chiqish", "admin paneldan chiqish"}:
        context.user_data.pop("admin_state", None)
        context.user_data.pop("admin_action", None)
        await update.message.reply_text("Admin panel yopildi ✅", reply_markup=ReplyKeyboardRemove())
        return ConversationHandler.END

    await update.message.reply_text("Iltimos, tugmadan tanlang: Test yaratish, Natijalarni olish yoki Test o'chirish.")
    return ADMIN_ACTION


async def admin_test_number(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    test_number = (update.message.text or "").strip()
    if not test_number.isdigit():
        await update.message.reply_text("Faqat raqam kiriting. Masalan: 55")
        return ADMIN_TEST_NUMBER

    mode = context.user_data.get("admin_action")

    if mode == "create":
        context.user_data["test_number"] = test_number
        context.user_data["admin_state"] = "create_keys"
        await update.message.reply_text("Kalitlarni kiriting (faqat a,b,c,d). Masalan: abbcd 🔐")
        return ADMIN_KEYS

    if mode == "results":
        excel_stream, result_info = build_results_excel(test_number)
        if excel_stream is not None:
            await update.message.reply_document(
                document=InputFile(excel_stream, filename=result_info),
                caption=f"📊 {test_number}-test natijalari (Excel)",
            )
            await update.message.reply_text(get_test_results_text(test_number))
        else:
            await update.message.reply_text(result_info)

        context.user_data["admin_state"] = "action"
        await update.message.reply_text("Yana amal tanlang 👇", reply_markup=admin_keyboard())
        return ADMIN_ACTION

    if mode == "delete":
        ok, msg = delete_test_by_number(test_number)
        await update.message.reply_text(msg)
        context.user_data["admin_state"] = "action"
        await update.message.reply_text("Yana amal tanlang 👇", reply_markup=admin_keyboard())
        return ADMIN_ACTION

    await update.message.reply_text("Iltimos, amalni boshidan tanlang.", reply_markup=admin_keyboard())
    return ADMIN_ACTION

async def admin_keys(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    keys = (update.message.text or "").strip().lower()
    test_number = context.user_data.get("test_number", "")

    ok, msg = create_test_from_keys(test_number, keys)
    await update.message.reply_text(msg)

    if ok:
        tg_id = update.effective_user.id
        user = get_user_by_telegram_id(tg_id)
        if user:
            test_title = f"{test_number}-test"
            access = None
            with db_conn() as conn:
                row = conn.execute("SELECT access_code FROM tests WHERE title = ?", (test_title,)).fetchone()
                access = (row["access_code"] if row else None)
            site_link = (
                f"{build_site_link(tg_id, user['access_key'], user['first_name'], user['last_name'], user['region'])}"
                f"&q={quote_plus(test_title)}&tc={quote_plus(access or '')}&v={int(datetime.now().timestamp())}"
            )
            await update.message.reply_text(f"Yaratilgan testni tekshirish: {site_link}")

    context.user_data["admin_state"] = "action"
    await update.message.reply_text("Yana amal tanlang 👇", reply_markup=admin_keyboard())
    return ADMIN_ACTION


async def image_test_title(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    title = (update.message.text or "").strip()
    if len(title) < 2:
        await update.message.reply_text("Test nomi kamida 2 ta belgidan iborat bo'lsin.")
        return IMG_TEST_TITLE

    with db_conn() as conn:
        exists = conn.execute("SELECT id FROM tests WHERE title = ?", (title,)).fetchone()
        if exists:
            await update.message.reply_text("Bu nomdagi test allaqachon mavjud. Boshqa nom kiriting.")
            return IMG_TEST_TITLE

        code = generate_test_code(conn)
        cur = conn.execute(
            "INSERT INTO tests (title, description, access_code) VALUES (?, ?, ?)",
            (title, "Rasmli test", code),
        )
        test_id = cur.lastrowid
        conn.commit()

    context.user_data["img_test_id"] = test_id
    context.user_data["img_test_title"] = title
    context.user_data["img_test_code"] = code
    context.user_data["img_question_index"] = 1
    context.user_data["admin_state"] = "image_photo"

    await update.message.reply_text("1-savol rasmini yuboring.")
    return IMG_QUESTION_IMAGE


async def image_test_photo(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    test_id = context.user_data.get("img_test_id")
    index = context.user_data.get("img_question_index", 1)
    if not test_id:
        await update.message.reply_text("Test topilmadi. /admin orqali qayta boshlang.")
        return ConversationHandler.END

    image_path = await save_telegram_photo(update, test_id, index)
    if not image_path:
        await update.message.reply_text("Iltimos, savol rasmini yuboring.")
        return IMG_QUESTION_IMAGE

    context.user_data["img_image_path"] = image_path
    context.user_data["admin_state"] = "image_answer"
    await update.message.reply_text("To'g'ri javobni kiriting (a/b/c/d):")
    return IMG_QUESTION_ANSWER


async def image_test_answer(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    answer = (update.message.text or "").strip().lower()
    if answer not in {"a", "b", "c", "d"}:
        await update.message.reply_text("Faqat a, b, c, d kiriting.")
        return IMG_QUESTION_ANSWER

    test_id = context.user_data.get("img_test_id")
    index = context.user_data.get("img_question_index", 1)
    image_path = context.user_data.get("img_image_path")
    if not test_id or not image_path:
        await update.message.reply_text("Rasm topilmadi. Qayta urinib ko'ring.")
        return IMG_QUESTION_IMAGE

    with db_conn() as conn:
        conn.execute(
            """
            INSERT INTO questions (test_id, text, option_a, option_b, option_c, option_d, image_path, correct_option)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (test_id, f"{index}-savol", "A", "B", "C", "D", image_path, answer),
        )
        conn.commit()

    context.user_data["img_question_index"] = index + 1
    context.user_data["admin_state"] = "image_next"
    await update.message.reply_text("Yana savol qo'shasizmi?", reply_markup=image_test_next_keyboard())
    return IMG_QUESTION_NEXT


async def image_test_next(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    text = (update.message.text or "").strip().lower()
    if "yana" in text:
        index = context.user_data.get("img_question_index", 1)
        context.user_data["admin_state"] = "image_photo"
        await update.message.reply_text(f"{index}-savol rasmini yuboring.", reply_markup=ReplyKeyboardRemove())
        return IMG_QUESTION_IMAGE

    if "bo'ldi" in text or "boldi" in text:
        tg_id = update.effective_user.id
        user = get_user_by_telegram_id(tg_id)
        test_title = context.user_data.get("img_test_title", "")
        code = context.user_data.get("img_test_code", "")
        if user and test_title and code:
            site_link = (
                f"{build_site_link(tg_id, user['access_key'], user['first_name'], user['last_name'], user['region'])}"
                f"&q={quote_plus(test_title)}&tc={quote_plus(code)}&v={int(datetime.now().timestamp())}"
            )
            await update.message.reply_text(f"Test yaratildi. Test kodi: {code}")
            await update.message.reply_text(f"Yaratilgan testni tekshirish: {site_link}")
        context.user_data.pop("img_test_id", None)
        context.user_data.pop("img_test_title", None)
        context.user_data.pop("img_test_code", None)
        context.user_data.pop("img_question_index", None)
        context.user_data.pop("img_image_path", None)
        context.user_data.pop("admin_state", None)
        return ConversationHandler.END

    await update.message.reply_text("Iltimos, 'Yana savol' yoki 'Bo'ldi' ni tanlang.", reply_markup=image_test_next_keyboard())
    return IMG_QUESTION_NEXT


async def admin_cancel(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data.pop("admin_state", None)
    context.user_data.pop("admin_action", None)
    await update.message.reply_text("Admin panel yopildi ✅", reply_markup=ReplyKeyboardRemove())
    return ConversationHandler.END


async def admin_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    await admin_panel(update, context)


async def admin_delete_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    tg_id = update.effective_user.id
    if not is_admin(tg_id):
        await update.message.reply_text("Kechirasiz, sizda admin huquqi yo'q ⛔")
        return

    if context.args and context.args[0].isdigit():
        ok, msg = delete_test_by_number(context.args[0])
        await update.message.reply_text(msg)
        return

    context.user_data["admin_action"] = "delete"
    context.user_data["admin_state"] = "delete_number"
    await update.message.reply_text("Qaysi testni o'chiramiz? Raqamini kiriting (masalan: 12) 🗑")


async def add_admin_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    tg_id = update.effective_user.id
    if not is_super_admin(tg_id):
        await update.message.reply_text("Kechirasiz, bu buyruq faqat asosiy admin uchun.")
        return

    context.user_data["admin_state"] = "add_admin"
    await update.message.reply_text("Admin qo'shmoqchimisiz? Unda uning Telegram ID sini tashlang.")


async def remove_admin_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    tg_id = update.effective_user.id
    if not is_super_admin(tg_id):
        await update.message.reply_text("Kechirasiz, bu buyruq faqat asosiy admin uchun.")
        return

    context.user_data["admin_state"] = "remove_admin"
    await update.message.reply_text("Qaysi adminni o'chiramiz? Telegram ID sini yuboring.")


async def list_admins_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    tg_id = update.effective_user.id
    if not is_super_admin(tg_id):
        await update.message.reply_text("Kechirasiz, bu buyruq faqat asosiy admin uchun.")
        return

    rows = list_admins()
    if not rows:
        await update.message.reply_text("Adminlar ro'yxati bo'sh.")
        return

    lines = ["Adminlar ro'yxati:"]
    for r in rows:
        label = "super" if r["is_super"] else "admin"
        lines.append(f"- {r['telegram_id']} ({label})")
    await update.message.reply_text("\n".join(lines))

async def admin_text_router(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    tg_id = update.effective_user.id
    if not is_admin(tg_id):
        return

    state = context.user_data.get("admin_state")
    if not state:
        return

    if state == "action":
        await admin_action(update, context)
        return
    if state in {"create_number", "results_number", "delete_number"}:
        await admin_test_number(update, context)
        return
    if state == "create_keys":
        await admin_keys(update, context)
        return
    if state == "image_title":
        await image_test_title(update, context)
        return
    if state == "image_photo":
        await image_test_photo(update, context)
        return
    if state == "image_answer":
        await image_test_answer(update, context)
        return
    if state == "image_next":
        await image_test_next(update, context)
        return
    if state == "add_admin":
        if not is_super_admin(tg_id):
            await update.message.reply_text("Kechirasiz, bu buyruq faqat asosiy admin uchun.")
            context.user_data.pop("admin_state", None)
            return
        raw = (update.message.text or "").strip()
        if not raw.isdigit():
            await update.message.reply_text("Telegram ID raqam bo'lishi kerak. Qayta kiriting.")
            return
        admin_id = int(raw)
        ok = add_admin(admin_id)
        await update.message.reply_text("Admin qo'shildi." if ok else "Admin qo'shilmadi.")
        context.user_data.pop("admin_state", None)
        return
    if state == "remove_admin":
        if not is_super_admin(tg_id):
            await update.message.reply_text("Kechirasiz, bu buyruq faqat asosiy admin uchun.")
            context.user_data.pop("admin_state", None)
            return
        raw = (update.message.text or "").strip()
        if not raw.isdigit():
            await update.message.reply_text("Telegram ID raqam bo'lishi kerak. Qayta kiriting.")
            return
        admin_id = int(raw)
        ok = remove_admin(admin_id)
        await update.message.reply_text("Admin o'chirildi." if ok else "Admin o'chirilmadi.")
        context.user_data.pop("admin_state", None)
        return


async def admin_photo_router(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    tg_id = update.effective_user.id
    if not is_admin(tg_id):
        return

    state = context.user_data.get("admin_state")
    if state == "image_photo":
        await image_test_photo(update, context)


async def on_error(update: object, context: ContextTypes.DEFAULT_TYPE) -> None:
    msg = str(context.error) if context.error else "Unknown error"
    if "terminated by other getUpdates request" in msg:
        print("[BOT ERROR] Conflict: bot token bilan boshqa instance ham ishlayapti.")
        print("[BOT ERROR] Boshqa bot jarayonlarini to'xtating va faqat bitta instance qoldiring.")
        return
    if "wrong http url" in msg:
        print("[BOT ERROR] Inline tugma URL noto'g'ri. BASE_SITE_URL public URL bo'lishi kerak.")
        return
    print(f"[BOT ERROR] {msg}")


def run_flask() -> None:
    app.run(host="0.0.0.0", port=5000, debug=False, use_reloader=False)


def run_bot() -> None:
    if BOT_TOKEN == "TOKENNI_BU_YERGA_QOYING":
        raise RuntimeError("BOT_TOKEN topilmadi. Uni .env fayliga yoki muhit o'zgaruvchisiga kiriting.")

    application = Application.builder().token(BOT_TOKEN).build()

    user_conv = ConversationHandler(
        entry_points=[CommandHandler("start", start)],
        states={
            REGISTER_FIRST: [MessageHandler(filters.TEXT & ~filters.COMMAND, first_name)],
            REGISTER_LAST: [MessageHandler(filters.TEXT & ~filters.COMMAND, last_name)],
            REGISTER_REGION: [MessageHandler(filters.TEXT & ~filters.COMMAND, region)],
            START_CHOICE: [MessageHandler(filters.TEXT & ~filters.COMMAND, start_choice)],
            TEST_CODE: [MessageHandler(filters.TEXT & ~filters.COMMAND, test_code)],
        },
        fallbacks=[CommandHandler("cancel", cancel), CommandHandler("cencel", cancel)],
        allow_reentry=True,
    )

    edit_conv = ConversationHandler(
        entry_points=[
            CommandHandler(["editprofile", "edit"], editprofile),
            CallbackQueryHandler(editprofile_callback, pattern=r"^edit_profile$"),
        ],
        states={
            EDIT_FIRST: [MessageHandler(filters.TEXT & ~filters.COMMAND, edit_first)],
            EDIT_LAST: [MessageHandler(filters.TEXT & ~filters.COMMAND, edit_last)],
            EDIT_REGION: [MessageHandler(filters.TEXT & ~filters.COMMAND, edit_region)],
        },
        fallbacks=[CommandHandler("cancel", cancel), CommandHandler("cencel", cancel)],
        allow_reentry=True,
    )
    application.add_handler(CommandHandler(["admin", "panel"], admin_command), group=-1)
    application.add_handler(CommandHandler("delete", admin_delete_command), group=-1)
    application.add_handler(CommandHandler("addadmin", add_admin_command), group=-1)
    application.add_handler(CommandHandler("removeadmin", remove_admin_command), group=-1)
    application.add_handler(CommandHandler("admins", list_admins_command), group=-1)
    application.add_handler(CommandHandler("canceladmin", admin_cancel), group=-1)
    application.add_handler(CommandHandler("cenceladmin", admin_cancel), group=-1)
    application.add_handler(CommandHandler("cencel", cancel), group=-1)
    application.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, admin_text_router), group=-1)
    application.add_handler(MessageHandler(filters.PHOTO, admin_photo_router), group=-1)

    application.add_handler(user_conv)
    application.add_handler(edit_conv)
    application.add_error_handler(on_error)

    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    application.run_polling()


if __name__ == "__main__":
    init_db()

    flask_thread = threading.Thread(target=run_flask, daemon=True)
    flask_thread.start()

    run_bot()





















































